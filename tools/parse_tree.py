"""Prototype converter: reads Hanh's visual family-tree workbook and derives a nested JSON tree.

Reading rules (PROTOTYPE, pending Hanh's confirmation; see SESSION_LEDGER Entry 2):
- A person = a ♂/♀ symbol cell with the name two columns to its right.
- Generation = which symbol column the person sits in (left-most = generation 1).
- Entries in the same column exactly 2 rows apart share one box (couple or a person + spouses).
- A box's parent box = the nearest box above it in the previous generation column.
- Angel-wings picture anchored on a name cell = deceased.
- Other text in a symbol column (e.g. "Any kids?") = a note on the box covering that row.
- Names are copied exactly. Nothing is merged across sheets.
"""
import json, sys
import openpyxl
from openpyxl.utils import column_index_from_string as C, get_column_letter as L

SHEETS = ["H.HN.T Paternal Tree", "H.HN.T Maternal Tree", "Huong's Paternal Tree", "Huong's Maternal Tree"]
SYM = {"♂": "M", "♀": "F"}

def parse(ws):
    wings = set()
    for im in ws._images:
        a = im.anchor._from
        wings.add((a.row + 1, a.col + 1))
    syms, texts = [], []
    for row in ws.iter_rows():
        for c in row:
            v = c.value
            if v is None or (isinstance(v, str) and not v.strip()):
                continue
            if isinstance(v, str) and v.strip() in SYM:
                syms.append((c.row, c.column, SYM[v.strip()]))
            else:
                texts.append((c.row, c.column, str(v)))
    gcols = sorted({col for _, col, _ in syms})
    textmap = {(r, col): t for r, col, t in texts}
    used = set()
    people = []
    for r, col, g in syms:
        name = textmap.get((r, col + 2))
        used.add((r, col + 2))
        people.append(dict(row=r, col=col, gen=gcols.index(col) + 1, gender=g,
                           name=name if name is not None else "",
                           deceased=(r, col + 2) in wings, cell=f"{L(col+2)}{r}"))
    # name-column text with no symbol (e.g. "Preggo")
    namecols = {c + 2 for c in gcols}
    notes = []
    for (r, col), t in textmap.items():
        if (r, col) in used:
            continue
        if col in namecols:
            people.append(dict(row=r, col=col - 2, gen=gcols.index(col - 2) + 1, gender=None, name=t,
                               deceased=False, cell=f"{L(col)}{r}", flag="no gender symbol"))
        else:
            notes.append(dict(row=r, col=col, text=t, cell=f"{L(col)}{r}"))
    # boxes
    boxes = []
    for gc in gcols:
        col_people = sorted([p for p in people if p["col"] == gc], key=lambda p: p["row"])
        cur = None
        for p in col_people:
            if cur and p["row"] - cur["members"][-1]["row"] == 2:
                cur["members"].append(p)
            else:
                cur = dict(gen=p["gen"], row=p["row"], members=[p], children=[], notes=[])
                boxes.append(cur)
    boxes.sort(key=lambda b: (b["row"], b["gen"]))
    roots = []
    for b in boxes:
        par = [x for x in boxes if x["gen"] == b["gen"] - 1 and x["row"] < b["row"]]
        if par:
            max(par, key=lambda x: x["row"])["children"].append(b)
        else:
            roots.append(b)
    for n in notes:  # attach note to the box whose span covers the row (nearest box starting at/above it)
        cand = [b for b in boxes if b["row"] <= n["row"]]
        if cand:
            tgt = max(cand, key=lambda b: (b["row"], b["gen"]))
            tgt["notes"].append(n["text"])
    for b in boxes:
        b["children"].sort(key=lambda x: x["row"])
    return roots, len(people), len(boxes), notes

def clean(b):
    return dict(gen=b["gen"], row=b["row"],
                members=[{k: m[k] for k in ("name", "gender", "deceased", "cell") if k in m} | ({"flag": m["flag"]} if "flag" in m else {}) for m in b["members"]],
                notes=b["notes"], children=[clean(c) for c in b["children"]])

def apply_corrections(tree_by_sheet, path):
    """Apply Hanh-approved spelling fixes (data/corrections.json). Each fix must match the sheet text exactly."""
    import os
    if not os.path.exists(path): return 0
    fixes = {k: v for k, v in json.load(open(path, encoding="utf-8")).items() if not k.startswith("_")}
    n = 0
    def walk(b, sheet):
        nonlocal n
        for m in b["members"]:
            k = sheet + "!" + m["cell"]
            if k in fixes:
                assert m["name"] == fixes[k]["from"], f"{k}: sheet has {m['name']!r}, fix expects {fixes[k]['from']!r}"
                m["name"] = fixes[k]["to"]; n += 1
        for c in b["children"]: walk(c, sheet)
    for s, v in tree_by_sheet.items():
        for r in v["roots"]: walk(r, s)
    assert n == len(fixes), f"applied {n} of {len(fixes)} corrections"
    return n

def apply_structure(tree_by_sheet, path):
    """Apply Hanh-confirmed structure fixes (data/structure.json): who is the blood relative in a box,
    moving children between boxes, and which wife each child belongs to. Boxes/children are named by the
    name cell of their first-listed person on the sheet. Every reference must resolve (asserted)."""
    import os
    if not os.path.exists(path): return 0
    spec = {k: v for k, v in json.load(open(path, encoding="utf-8")).items() if not k.startswith("_")}
    n = 0
    for sheet, ops in spec.items():
        roots = tree_by_sheet[sheet]["roots"]
        def find(cell):
            st = list(roots)
            while st:
                b = st.pop()
                if any(m["cell"] == cell for m in b["members"]): return b
                st += b["children"]
            raise AssertionError(f"{sheet}: no box with {cell}")
        for op in ops:
            b = find(op["box"])
            if op["op"] == "blood_first":
                i = [m["cell"] for m in b["members"]].index(op["person"])
                b["members"].insert(0, b["members"].pop(i))
            elif op["op"] == "move_children":
                src = find(op["from"])
                moved = []
                for c in op["children"]:
                    k = [x for x in src["children"] if x["members"][0]["cell"] == c]
                    assert len(k) == 1, f"{sheet}: {c} is not a child of {op['from']}"
                    src["children"].remove(k[0]); moved.append(k[0])
                pos = op.get("position", len(b["children"]))
                b["children"][pos:pos] = moved
            elif op["op"] == "mothers":
                cells = {m["cell"] for m in b["members"][1:]}
                kids = {x["members"][0]["cell"]: x for x in b["children"]}
                assert set(op["children"]) == set(kids), f"{sheet}: mothers must cover every child of {op['box']}"
                for c, mom in op["children"].items():
                    assert mom in cells, f"{sheet}: {mom} is not a spouse in {op['box']}"
                    kids[c]["mother"] = mom
            elif op["op"] == "gender":   # the sheet's ♂/♀ was wrong (Hanh confirmed)
                m = [x for x in b["members"] if x["cell"] == op["person"]]
                assert len(m) == 1 and m[0]["gender"] == op["from"], f"{sheet}: {op['person']} gender is not {op['from']}"
                m[0]["gender"] = op["to"]
            elif op["op"] == "stepchild":   # data only, nothing drawn (Hanh, Entry 30)
                k = [x for x in b["children"] if x["members"][0]["cell"] == op["child"]]
                assert len(k) == 1, f"{sheet}: {op['child']} is not a child of {op['box']}"
                k[0]["step"] = op["note"]
            else:
                raise AssertionError("unknown op " + op["op"])
            n += 1
    return n

if __name__ == "__main__":
    wb = openpyxl.load_workbook(sys.argv[1])
    out = {}
    for s in SHEETS:
        roots, np_, nb, notes = parse(wb[s])
        out[s] = dict(people=np_, boxes=nb, roots=[clean(r) for r in roots])
        print(f"{s}: {np_} people, {nb} boxes, {len(roots)} root box(es), {len(notes)} notes", file=sys.stderr)
    import os
    corr = sys.argv[3] if len(sys.argv) > 3 else os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "data", "corrections.json")
    print(f"corrections applied: {apply_corrections(out, corr)}", file=sys.stderr)
    struct = os.path.join(os.path.dirname(os.path.abspath(corr)), "structure.json")
    print(f"structure fixes applied: {apply_structure(out, struct)}", file=sys.stderr)
    json.dump(out, open(sys.argv[2], "w", encoding="utf-8"), ensure_ascii=False, indent=1)
